from __future__ import annotations

import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from typing import Any
from unittest.mock import patch

from openpyxl import Workbook

from app.backend import DatabaseManager
from app.drive_import import DriveImportError


def _workbook_bytes(sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> bytes:
    workbook = Workbook()
    first_sheet = True
    for sheet_name, (headers, rows) in sheets.items():
        if first_sheet:
            sheet = workbook.active
            sheet.title = sheet_name
            first_sheet = False
        else:
            sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _source_workbooks() -> dict[str, bytes]:
    return {
        "Applications": _workbook_bytes(
            {
                "Applications": (
                    [
                        "Timestamp",
                        "Full Name",
                        "Desired Training",
                        "Province",
                        "Current Status",
                        "Processing Status",
                        "Email",
                    ],
                    [
                        [
                            "2023-04-01",
                            "Applicant One",
                            "Software Testing",
                            "Utrecht",
                            "In the camp, waiting for residence-permit interview",
                            "DONE",
                            "one@example.invalid",
                        ],
                        [
                            "2023-04-02",
                            "Applicant Two",
                            "Data Engineer",
                            "Flevoland",
                            "Moved into housing (studying the language)",
                            None,
                            "two@example.invalid",
                        ],
                    ],
                )
            }
        ),
        "Interviews": _workbook_bytes(
            {
                "Interviews": (
                    ["Full Name", "Project Sent Date", "Project Received Date"],
                    [["Applicant One", "2023-05-31", "2023-08-07"]],
                )
            }
        ),
        "Mentor": _workbook_bytes(
            {
                "Mentor": (
                    ["Date", "Candidate Name", "Mentor Name", "Score", "Notes"],
                    [["2023-05-19", "Applicant Two", "Mentor One", 8, "Review note"]],
                ),
                "Mentor Recommendation Options": (
                    ["Recommendation Option (English)"],
                    [["Continue"]],
                ),
            }
        ),
        "Users": _workbook_bytes(
            {
                "Users": (
                    ["Username", "Password", "Role"],
                    [["admin.test", "Never-print-this", "admin"], ["resident.test", "Other-secret", "user"]],
                )
            }
        ),
    }


class DriveImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temporary_directory.name) / "test.db"
        self.database = DatabaseManager(self.database_path, import_drive_data=False)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_imports_source_rows_once_and_hashes_user_passwords(self) -> None:
        sources = _source_workbooks()

        self.assertTrue(self.database.import_workbooks(sources))
        self.assertFalse(self.database.import_workbooks(sources))

        applications = self.database.list_applications()
        self.assertEqual(len(applications), 2)
        self.assertEqual(applications[0]["current_status"], "Moved into housing (studying the language)")
        self.assertEqual(applications[0]["status"], "Not provided")
        self.assertEqual(applications[0]["priority"], "Not specified")

        residents = self.database.list_residents()
        self.assertEqual(residents, [])

        meetings = self.database.list_meetings()
        self.assertEqual(len(meetings), 1)
        self.assertEqual(meetings[0]["meeting_time"], "2023-05-19")
        self.assertEqual(meetings[0]["status"], "Completed")

        application_records = self.database.list_source_records(
            "Applications", "Applications"
        )
        self.assertEqual(len(application_records), 2)
        self.assertEqual(
            application_records[0]["headers"],
            [
                "Timestamp",
                "Full Name",
                "Desired Training",
                "Province",
                "Current Status",
                "Processing Status",
                "Email",
            ],
        )
        interview_records = self.database.list_source_records(
            "Interviews", "Interviews"
        )
        self.assertEqual(
            interview_records[0]["values"],
            ["Applicant One", "2023-05-31", "2023-08-07"],
        )
        mentor_records = self.database.list_source_records("Mentor", "Mentor")
        self.assertEqual(mentor_records[0]["values"][2:], ["Mentor One", 8, "Review note"])

        user = self.database.get_authenticated_user("admin.test", "Never-print-this")
        self.assertEqual(user, {"username": "admin.test", "role": "admin", "display_name": "admin.test"})
        self.assertIsNone(self.database.get_authenticated_user("admin.test", "wrong"))
        regular_user = self.database.get_authenticated_user(
            "resident.test", "Other-secret"
        )
        self.assertEqual(regular_user["role"], "user")

        with self.database._connect() as connection:
            stored_password = connection.execute(
                "SELECT password FROM users WHERE username = 'admin.test'"
            ).fetchone()[0]
            source_rows = [
                row[0]
                for row in connection.execute(
                    "SELECT payload_json FROM source_records WHERE workbook_name = 'Users'"
                ).fetchall()
            ]
            marker_count = connection.execute(
                "SELECT COUNT(*) FROM app_meta WHERE key = 'drive_import_completed'"
            ).fetchone()[0]
        self.assertTrue(stored_password.startswith("pbkdf2_sha256$"))
        self.assertNotIn("Never-print-this", json.dumps(source_rows))
        self.assertNotIn("Other-secret", json.dumps(source_rows))
        self.assertEqual(marker_count, 1)
        self.assertEqual(len(source_rows), 2)
        with patch("app.backend.download_workbooks") as download:
            DatabaseManager(self.database_path)
        download.assert_not_called()

    def test_missing_workbook_does_not_change_database(self) -> None:
        sources = _source_workbooks()
        del sources["Users"]

        with self.assertRaises(DriveImportError):
            self.database.import_workbooks(sources)

        self.assertEqual(self.database.list_applications(), [])
        with self.database._connect() as connection:
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM source_records").fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM app_meta").fetchone()[0],
                0,
            )

    def test_candidate_source_records_can_be_added_and_edited(self) -> None:
        self.database.import_workbooks(_source_workbooks())
        original = self.database.list_source_records(
            "Applications", "Applications"
        )[0]
        values = dict(zip(original["headers"], original["values"]))
        values["Full Name"] = "Applicant Updated"
        values["Email"] = "updated@example.invalid"
        self.database.save_application_source_record(
            values, source_row=original["source_row"]
        )

        new_values = {header: "" for header in original["headers"]}
        new_values.update(
            {
                "Full Name": "New Applicant",
                "Email": "new@example.invalid",
                "Application Period": "VIT3",
            }
        )
        new_row = self.database.save_application_source_record(new_values)
        records = self.database.list_source_records(
            "Applications", "Applications"
        )
        edited = next(
            record for record in records if record["source_row"] == original["source_row"]
        )
        added = next(record for record in records if record["source_row"] == new_row)
        self.assertEqual(
            dict(zip(edited["headers"], edited["values"]))["Email"],
            "updated@example.invalid",
        )
        self.assertEqual(
            dict(zip(added["headers"], added["values"]))["Full Name"],
            "New Applicant",
        )

    def test_corrupt_workbook_does_not_change_database(self) -> None:
        sources = _source_workbooks()
        sources["Mentor"] = b"not an Excel workbook"

        with self.assertRaises(DriveImportError):
            self.database.import_workbooks(sources)

        self.assertEqual(self.database.list_applications(), [])
        with self.database._connect() as connection:
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM source_records").fetchone()[0],
                0,
            )

    def test_invalid_user_row_rolls_back_entire_import(self) -> None:
        sources = _source_workbooks()
        sources["Users"] = _workbook_bytes(
            {
                "Users": (
                    ["Username", "Password", "Role"],
                    [["admin.test", "secret", "admin"], ["bad.user", "", "unsupported"]],
                )
            }
        )

        with self.assertRaises(DriveImportError):
            self.database.import_workbooks(sources)

        self.assertEqual(self.database.list_applications(), [])
        with self.database._connect() as connection:
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM source_records").fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM users").fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM app_meta").fetchone()[0],
                0,
            )

    def test_existing_plaintext_password_is_upgraded_after_successful_login(self) -> None:
        with self.database._connect() as connection:
            connection.execute(
                """
                INSERT INTO users (username, password, role, display_name)
                VALUES ('legacy', 'legacy-secret', 'resident', 'Legacy User')
                """
            )

        self.assertEqual(
            self.database.get_authenticated_user("legacy", "legacy-secret"),
            {"username": "legacy", "role": "resident", "display_name": "Legacy User"},
        )
        with self.database._connect() as connection:
            stored_password = connection.execute(
                "SELECT password FROM users WHERE username = 'legacy'"
            ).fetchone()[0]
        self.assertTrue(stored_password.startswith("pbkdf2_sha256$"))
        self.assertNotEqual(stored_password, "legacy-secret")


if __name__ == "__main__":
    unittest.main()
