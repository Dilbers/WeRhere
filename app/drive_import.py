from __future__ import annotations

from io import BytesIO
import re
from pathlib import Path
from typing import Any, Mapping
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException


DRIVE_SPREADSHEET_IDS = {
    "Applications": "1URhoWeH8U9o1d8gXvk_DU4kkGHMcpTF0RI4Ey5Inngg",
    "Interviews": "1lMQwEFr39dmDeSy5Tv1AJrYCc-QIZHFe9oByKYy1Q3o",
    "Mentor": "1yWwAkgRlNsPUwNb9ZPKq5SnC3uef0oZMmyZ7_u67jpE",
    "Users": "1zjxlJJIRVOYPVLGX-1FZdxfOpyozxO3MZgw5i66OPPU",
}

PROJECT_ROOT = Path(__file__).resolve().parent.parent
GOOGLE_CREDENTIALS_DIR = PROJECT_ROOT / "credentials"
GOOGLE_CLIENT_FILE = GOOGLE_CREDENTIALS_DIR / "google_oauth_client.json"
GOOGLE_TOKEN_FILE = GOOGLE_CREDENTIALS_DIR / "google_oauth_token.json"
GOOGLE_SHEETS_SCOPE = "https://www.googleapis.com/auth/spreadsheets"

REQUIRED_SHEETS = {
    "Applications": {
        "Applications": {"Full Name", "Desired Training", "Province", "Current Status", "Processing Status", "Timestamp"},
    },
    "Interviews": {
        "Interviews": {"Full Name", "Project Sent Date", "Project Received Date"},
    },
    "Mentor": {
        "Mentor": {"Date", "Candidate Name", "Mentor Name"},
        "Mentor Recommendation Options": {"Recommendation Option (English)"},
    },
    "Users": {
        "Users": {"Username", "Password", "Role"},
    },
}

class DriveImportError(RuntimeError):
    """Raised when the shared source data cannot be downloaded or validated."""


class GoogleSheetsStore:
    """Authenticated read/write access to the user's copied Sheets workbooks."""

    def __init__(self):
        self._service = None

    def _get_service(self):
        if self._service is not None:
            return self._service
        try:
            from google.auth.transport.requests import Request as GoogleAuthRequest
            from google.oauth2.credentials import Credentials
            from google_auth_oauthlib.flow import InstalledAppFlow
            from googleapiclient.discovery import build
            from googleapiclient.errors import HttpError
        except ImportError as error:
            raise DriveImportError(
                "Google Sheets support is not installed. Install the packages in requirements.txt."
            ) from error

        if not GOOGLE_CLIENT_FILE.is_file():
            raise DriveImportError(
                "Google OAuth is not configured. Put the downloaded Desktop OAuth client JSON at "
                f"{GOOGLE_CLIENT_FILE}."
            )

        credentials = None
        if GOOGLE_TOKEN_FILE.is_file():
            try:
                credentials = Credentials.from_authorized_user_file(
                    str(GOOGLE_TOKEN_FILE), [GOOGLE_SHEETS_SCOPE]
                )
            except (OSError, ValueError) as error:
                raise DriveImportError(
                    f"The saved Google authorization is unreadable: {error}"
                ) from error

        if credentials and credentials.expired and credentials.refresh_token:
            try:
                credentials.refresh(GoogleAuthRequest())
            except Exception:
                credentials = None
        if not credentials or not credentials.valid:
            try:
                flow = InstalledAppFlow.from_client_secrets_file(
                    str(GOOGLE_CLIENT_FILE), [GOOGLE_SHEETS_SCOPE]
                )
                credentials = flow.run_local_server(port=0, open_browser=True)
            except Exception as error:
                raise DriveImportError(
                    f"Google sign-in did not complete: {error}"
                ) from error

        GOOGLE_CREDENTIALS_DIR.mkdir(parents=True, exist_ok=True)
        GOOGLE_TOKEN_FILE.write_text(credentials.to_json(), encoding="utf-8")
        try:
            self._service = build(
                "sheets", "v4", credentials=credentials, cache_discovery=False
            )
        except Exception as error:
            if isinstance(error, HttpError):
                raise DriveImportError(f"Could not connect to Google Sheets: {error}") from error
            raise
        return self._service

    def download_workbooks(self) -> dict[str, bytes]:
        service = self._get_service()
        workbooks: dict[str, bytes] = {}
        for workbook_name, spreadsheet_id in DRIVE_SPREADSHEET_IDS.items():
            required = REQUIRED_SHEETS[workbook_name]
            try:
                spreadsheet = (
                    service.spreadsheets()
                    .get(
                        spreadsheetId=spreadsheet_id,
                        fields="sheets.properties(title,gridProperties(rowCount,columnCount))",
                    )
                    .execute()
                )
                sheet_properties = {
                    sheet["properties"]["title"]: sheet["properties"]
                    for sheet in spreadsheet.get("sheets", [])
                }
                missing = set(required) - set(sheet_properties)
                if missing:
                    raise DriveImportError(
                        f"The {workbook_name} spreadsheet is missing required tab(s): "
                        f"{', '.join(sorted(missing))}."
                    )
                ranges = []
                for sheet_name in required:
                    grid = sheet_properties[sheet_name].get("gridProperties", {})
                    last_column = get_column_letter(max(1, grid.get("columnCount", 26)))
                    last_row = max(1, grid.get("rowCount", 1000))
                    ranges.append(f"'{sheet_name}'!A1:{last_column}{last_row}")
                result = (
                    service.spreadsheets()
                    .values()
                    .batchGet(
                        spreadsheetId=spreadsheet_id,
                        ranges=ranges,
                        valueRenderOption="FORMATTED_VALUE",
                    )
                    .execute()
                )
            except Exception as error:
                raise DriveImportError(
                    f"Could not read the {workbook_name} Google Sheet. Check its ID, tab names, "
                    f"and that the Google account has access. {error}"
                ) from error

            workbook = Workbook()
            workbook.remove(workbook.active)
            for value_range in result.get("valueRanges", []):
                sheet_title = value_range.get("range", "").split("!", 1)[0]
                sheet_title = sheet_title.strip("'").replace("''", "'")
                if sheet_title not in required:
                    continue
                rows = value_range.get("values", [])
                if not rows:
                    raise DriveImportError(
                        f"The {workbook_name} sheet tab '{sheet_title}' is empty."
                    )
                sheet = workbook.create_sheet(sheet_title)
                for row_number, row in enumerate(rows, start=1):
                    for column_number, value in enumerate(row, start=1):
                        sheet.cell(row=row_number, column=column_number, value=value)

            missing = set(required) - set(workbook.sheetnames)
            if missing:
                raise DriveImportError(
                    f"The {workbook_name} spreadsheet is missing required tab(s): "
                    f"{', '.join(sorted(missing))}."
                )
            buffer = BytesIO()
            workbook.save(buffer)
            workbook.close()
            workbooks[workbook_name] = buffer.getvalue()
        return workbooks

    def append_application(self, headers: list[str], values: Mapping[str, Any]) -> int:
        row = [[values.get(header, "") for header in headers]]
        end_column = get_column_letter(len(headers))
        try:
            response = (
                self._get_service()
                .spreadsheets()
                .values()
                .append(
                    spreadsheetId=DRIVE_SPREADSHEET_IDS["Applications"],
                    range=f"'Applications'!A:{end_column}",
                    valueInputOption="RAW",
                    insertDataOption="INSERT_ROWS",
                    body={"values": row},
                )
                .execute()
            )
        except Exception as error:
            raise DriveImportError(f"Could not add the applicant to Google Drive: {error}") from error
        updated_range = response.get("updates", {}).get("updatedRange", "")
        match = re.search(r"!A(\d+)", updated_range)
        if not match:
            raise DriveImportError(
                "Google Sheets accepted the applicant but did not return its row number. "
                "Refresh the Applications table before editing it."
            )
        return int(match.group(1))

    def append_mentor_record(self, headers: list[str], values: Mapping[str, Any]) -> int:
        row = [[values.get(header, "") for header in headers]]
        end_column = get_column_letter(len(headers))
        try:
            response = (
                self._get_service()
                .spreadsheets()
                .values()
                .append(
                    spreadsheetId=DRIVE_SPREADSHEET_IDS["Mentor"],
                    range=f"'Mentor'!A:{end_column}",
                    valueInputOption="RAW",
                    insertDataOption="INSERT_ROWS",
                    body={"values": row},
                )
                .execute()
            )
        except Exception as error:
            raise DriveImportError(
                f"Could not add the mentor interview to Google Drive: {error}"
            ) from error
        updated_range = response.get("updates", {}).get("updatedRange", "")
        match = re.search(r"!A(\d+)", updated_range)
        if not match:
            raise DriveImportError(
                "Google Sheets accepted the interview but did not return its row number. "
                "Refresh the Mentor interviews table before recording another result."
            )
        return int(match.group(1))

    def update_application(
        self, headers: list[str], values: Mapping[str, Any], source_row: int
    ) -> None:
        end_column = get_column_letter(len(headers))
        try:
            (
                self._get_service()
                .spreadsheets()
                .values()
                .update(
                    spreadsheetId=DRIVE_SPREADSHEET_IDS["Applications"],
                    range=f"'Applications'!A{source_row}:{end_column}{source_row}",
                    valueInputOption="RAW",
                    body={"values": [[values.get(header, "") for header in headers]]},
                )
                .execute()
            )
        except Exception as error:
            raise DriveImportError(f"Could not update the applicant in Google Drive: {error}") from error


def download_workbooks() -> dict[str, bytes]:
    return GoogleSheetsStore().download_workbooks()


def read_workbooks(
    workbook_sources: Mapping[str, bytes],
) -> dict[str, list[dict[str, Any]]]:
    if set(workbook_sources) != set(REQUIRED_SHEETS):
        raise DriveImportError(
            "The shared folder must contain Applications.xlsx, Interviews.xlsx, "
            "Mentor.xlsx, and Users.xlsx."
        )

    parsed: dict[str, list[dict[str, Any]]] = {}
    for workbook_name, required_sheets in REQUIRED_SHEETS.items():
        try:
            workbook = load_workbook(
                BytesIO(workbook_sources[workbook_name]),
                read_only=True,
                data_only=False,
            )
        except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
            raise DriveImportError(
                f"{workbook_name}.xlsx is not a readable Excel workbook."
            ) from error

        try:
            if not set(required_sheets).issubset(workbook.sheetnames):
                raise DriveImportError(
                    f"{workbook_name}.xlsx is missing a required worksheet."
                )
            sheets = []
            for sheet_name, required_headers in required_sheets.items():
                sheet = workbook[sheet_name]
                iterator = sheet.iter_rows(values_only=True)
                first_row = next(iterator, None)
                if first_row is None:
                    raise DriveImportError(
                        f"{workbook_name}.xlsx has an empty {sheet_name} worksheet."
                    )
                headers = [
                    str(value).strip() if value is not None else ""
                    for value in first_row
                ]
                if any(headers.count(header) != 1 for header in required_headers):
                    raise DriveImportError(
                        f"{workbook_name}.xlsx has an unexpected {sheet_name} header row."
                    )
                rows = [
                    (row_number, list(row))
                    for row_number, row in enumerate(iterator, start=2)
                    if any(value is not None for value in row)
                ]
                sheets.append(
                    {
                        "sheet_name": sheet_name,
                        "headers": headers,
                        "rows": rows,
                    }
                )
            parsed[workbook_name] = sheets
        finally:
            workbook.close()
    return parsed
